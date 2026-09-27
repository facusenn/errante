"""Panel admin protegido por ADMIN_PASSWORD.

Incluye: dashboard de reservas, gestión de portfolio, presupuestos,
configuración del estimador y export a Excel. La confirmación de un turno
dispara (si están configuradas) las integraciones de Google Calendar y
WhatsApp — si no lo están, la app sigue andando con fallbacks.
"""
import io
import os
import time
from datetime import datetime

from flask import (Blueprint, current_app, flash, redirect, render_template,
                   request, send_file, session, url_for)
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from werkzeug.security import check_password_hash

from database import db
from models import PortfolioImage, QuoteRequest, Reserva
from services import google_calendar, whatsapp
from services import images
from services.estimator import CONFIG_GROUPS, get_config, set_config, seed_defaults

admin_bp = Blueprint('admin', __name__, url_prefix='/admin')

MAX_LOGIN_FAILURES = 5
LOCKOUT_SECONDS = 300


# ─────────────────────────────────────────────
# AUTH
# ─────────────────────────────────────────────

def _admin_password():
    return os.environ.get('ADMIN_PASSWORD', '')


def _password_is_hash(password):
    return password.startswith(('pbkdf2:', 'scrypt:'))


def _check_password(candidate):
    stored = _admin_password()
    if not stored:
        return False
    if candidate is None:
        return False
    if _password_is_hash(stored):
        try:
            return check_password_hash(stored, candidate)
        except ValueError:
            return False
    return candidate == stored


@admin_bp.before_request
def require_login():
    if request.endpoint == 'admin.login':
        return None
    if not session.get('admin_authenticated'):
        return redirect(url_for('admin.login'))
    return None


@admin_bp.route('/login', methods=['GET', 'POST'])
def login():
    if not _admin_password():
        flash('El panel está bloqueado: configurá ADMIN_PASSWORD en tu archivo .env.', 'error')
        return render_template('admin_login.html', configured=False)

    lock_until = session.get('admin_lock_until', 0)
    if lock_until and lock_until > time.time():
        restantes = int(lock_until - time.time())
        flash(f'Demasiados intentos fallidos. Esperá {restantes} segundos.', 'error')
        return render_template('admin_login.html', configured=True)

    if request.method == 'POST':
        candidate = request.form.get('password', '')
        if _check_password(candidate):
            session['admin_authenticated'] = True
            session.pop('admin_failures', None)
            session.pop('admin_lock_until', None)
            return redirect(url_for('admin.dashboard'))
        failures = session.get('admin_failures', 0) + 1
        session['admin_failures'] = failures
        if failures >= MAX_LOGIN_FAILURES:
            session['admin_lock_until'] = time.time() + LOCKOUT_SECONDS
            session.pop('admin_failures', None)
            flash('Demasiados intentos fallidos. Esperá 5 minutos.', 'error')
        else:
            flash('Contraseña incorrecta.', 'error')
    return render_template('admin_login.html', configured=True)


@admin_bp.route('/logout', methods=['POST'])
def logout():
    session.pop('admin_authenticated', None)
    flash('Sesión cerrada.', 'info')
    return redirect(url_for('public.index'))


# ─────────────────────────────────────────────
# DASHBOARD DE RESERVAS
# ─────────────────────────────────────────────

@admin_bp.route('/')
def dashboard():
    estado = request.args.get('estado', 'todos')
    query = Reserva.query.order_by(Reserva.fecha.asc())
    if estado != 'todos':
        query = query.filter_by(estado=estado)
    reservas = query.all()

    stats = {
        'total': Reserva.query.count(),
        'pendientes': Reserva.query.filter_by(estado='pendiente').count(),
        'confirmados': Reserva.query.filter_by(estado='confirmado').count(),
        'cancelados': Reserva.query.filter_by(estado='cancelado').count(),
        'quotes_total': QuoteRequest.query.count(),
        'quotes_pendientes': QuoteRequest.query.filter_by(status='pendiente').count(),
        'quotes_aprobadas': QuoteRequest.query.filter_by(status='aprobado').count(),
        'quotes_agendadas': QuoteRequest.query.filter_by(status='agendado').count(),
    }
    integrations = {
        'whatsapp': whatsapp.is_configured(),
        'calendar': google_calendar.is_configured(),
    }
    return render_template('admin.html', reservas=reservas, stats=stats,
                           estado_filtro=estado, integrations=integrations)


# ─────────────────────────────────────────────
# RESERVAS: estado, eliminación y confirmación
# ─────────────────────────────────────────────

def _confirmar_reserva(reserva):
    """Dispara WhatsApp y Google Calendar si están configuradas.

    Nunca levanta excepción: los fallos quedan registrados en el estado
    de la reserva para que el panel los muestre sin romper el flujo.
    """
    # WhatsApp
    try:
        if whatsapp.is_configured():
            reserva.whatsapp_status = 'enviado' if whatsapp.send_confirmation(reserva) else 'error'
        else:
            reserva.whatsapp_status = 'wa_me'
    except Exception:
        reserva.whatsapp_status = 'error'

    # Google Calendar
    try:
        if google_calendar.is_configured():
            available = google_calendar.check_availability(reserva.fecha, reserva.hora)
            reserva.availability_checked = available is not None
            if available is False:
                reserva.calendar_status = 'conflicto'
            else:
                event_id = google_calendar.create_event(reserva)
                reserva.calendar_event_id = event_id
                reserva.calendar_synced_at = datetime.utcnow()
                reserva.calendar_status = 'sincronizado'
        else:
            reserva.calendar_status = 'pendiente'
    except Exception:
        reserva.calendar_status = 'error'


@admin_bp.route('/reserva/<int:id>/estado', methods=['POST'])
def cambiar_estado(id):
    reserva = Reserva.query.get_or_404(id)
    nuevo_estado = request.form['estado']
    if nuevo_estado in ['pendiente', 'confirmado', 'cancelado']:
        reserva.estado = nuevo_estado
        if nuevo_estado == 'confirmado':
            _confirmar_reserva(reserva)
        db.session.commit()

        avisos = []
        if reserva.whatsapp_status == 'enviado':
            avisos.append('WhatsApp enviado')
        elif reserva.whatsapp_status == 'wa_me':
            avisos.append('WhatsApp no configurado (link wa.me disponible)')
        elif reserva.whatsapp_status == 'error':
            avisos.append('No se pudo enviar el WhatsApp')
        if reserva.calendar_status == 'sincronizado':
            avisos.append('evento creado en Google Calendar')
        elif reserva.calendar_status == 'conflicto':
            avisos.append('⚠ el horario figura ocupado en Google Calendar')
        elif reserva.calendar_status == 'error':
            avisos.append('falló el sync con Google Calendar')

        detalle = f' ({", ".join(avisos)})' if avisos and nuevo_estado == 'confirmado' else ''
        flash(f'Estado actualizado a "{nuevo_estado}".{detalle}', 'success')
    return redirect(url_for('admin.dashboard'))


@admin_bp.route('/reserva/<int:id>/eliminar', methods=['POST'])
def eliminar_reserva(id):
    reserva = Reserva.query.get_or_404(id)
    db.session.delete(reserva)
    db.session.commit()
    flash('Reserva eliminada.', 'info')
    return redirect(url_for('admin.dashboard'))


# ─────────────────────────────────────────────
# EXPORTAR EXCEL
# ─────────────────────────────────────────────

@admin_bp.route('/exportar-excel')
def exportar_excel():
    import openpyxl

    reservas = Reserva.query.order_by(Reserva.fecha.asc()).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Turnos'

    # Estilos
    header_fill = PatternFill('solid', start_color='1a1a2e', end_color='1a1a2e')
    header_font = Font(bold=True, color='e94560', size=11, name='Arial')
    alt_fill = PatternFill('solid', start_color='f8f4f0', end_color='f8f4f0')
    border = Border(bottom=Side(style='thin', color='dddddd'))
    center = Alignment(horizontal='center', vertical='center')

    # Título
    ws.merge_cells('A1:M1')
    titulo = ws['A1']
    titulo.value = '📋 Registro de Turnos — Estudio de Tatuajes'
    titulo.font = Font(bold=True, size=14, color='1a1a2e', name='Arial')
    titulo.alignment = Alignment(horizontal='center', vertical='center')
    titulo.fill = PatternFill('solid', start_color='e94560', end_color='e94560')
    ws.row_dimensions[1].height = 32

    ws.append([])

    # Encabezados
    encabezados = ['ID', 'Nombre', 'Email', 'Teléfono', 'Fecha', 'Hora',
                   'Estilo', 'Zona', 'Tamaño', '1ra vez', 'Estado', 'Descripción', 'Presup. #']
    ws.append(encabezados)
    for cell in ws[3]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
    ws.row_dimensions[3].height = 22

    # Datos
    estado_colores = {
        'confirmado': 'd4edda',
        'cancelado': 'f8d7da',
        'pendiente': 'fff3cd',
    }
    for i, r in enumerate(reservas, start=4):
        fila = [
            r.id, r.nombre, r.email, r.telefono,
            r.fecha.strftime('%d/%m/%Y'), r.hora,
            r.estilo, r.zona, r.tamano,
            'Sí' if r.es_primera_vez else 'No',
            r.estado.capitalize(), r.descripcion or '', r.quote_id or '',
        ]

        ws.append(fila)
        color = estado_colores.get(r.estado, 'ffffff')
        row_fill = PatternFill('solid', start_color=color, end_color=color)
        for cell in ws[i]:
            cell.fill = row_fill if r.estado != 'pendiente' else (alt_fill if i % 2 == 0 else PatternFill())
            cell.border = border
            cell.alignment = Alignment(vertical='center', wrap_text=True)
        ws.row_dimensions[i].height = 18

    anchos = [6, 22, 28, 16, 12, 8, 16, 16, 12, 12, 14, 40, 10]
    for idx, ancho in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(idx)].width = ancho

    # Hoja de presupuestos
    quotes = QuoteRequest.query.order_by(QuoteRequest.created_at.asc()).all()
    ws2 = wb.create_sheet('Presupuestos')
    ws2['A1'] = 'Presupuestos recibidos'
    ws2['A1'].font = Font(bold=True, size=13, name='Arial')
    ws2.append([])
    encabezados_q = ['ID', 'Fecha', 'Nombre', 'Email', 'Teléfono',
                     'Tamaño', 'Zona', 'Estimado (ARS)', 'Estado', 'Descripción']
    ws2.append(encabezados_q)
    for cell in ws2[3]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center

    for i, q in enumerate(quotes, start=4):
        ws2.append([
            q.id,
            q.created_at.strftime('%d/%m/%Y') if q.created_at else '',
            q.nombre, q.email, q.telefono,
            q.size, q.zone,
            q.estimated_price or 0,
            q.status.capitalize(),
            q.description or '',
        ])
        for cell in ws2[i]:
            cell.border = border
            cell.alignment = Alignment(vertical='center', wrap_text=True)

    anchos_q = [6, 12, 22, 28, 16, 14, 16, 14, 12, 40]
    for idx, ancho in enumerate(anchos_q, 1):
        ws2.column_dimensions[get_column_letter(idx)].width = ancho

    # Hoja resumen
    ws3 = wb.create_sheet('Resumen')
    ws3['A1'] = 'Resumen'
    ws3['A1'].font = Font(bold=True, size=13, color='ffffff', name='Arial')
    ws3['A1'].fill = PatternFill('solid', start_color='e94560', end_color='e94560')
    ws3.append(['Estado', 'Cantidad'])
    for cell in ws3[2]:
        cell.font = header_font
        cell.fill = header_fill

    filas_resumen = [
        ['Turnos pendientes', sum(1 for r in reservas if r.estado == 'pendiente')],
        ['Turnos confirmados', sum(1 for r in reservas if r.estado == 'confirmado')],
        ['Turnos cancelados', sum(1 for r in reservas if r.estado == 'cancelado')],
        ['Presupuestos pendientes', sum(1 for q in quotes if q.status == 'pendiente')],
        ['Presupuestos aprobados', sum(1 for q in quotes if q.status == 'aprobado')],
        ['Presupuestos agendados', sum(1 for q in quotes if q.status == 'agendado')],
        ['Presupuestos rechazados', sum(1 for q in quotes if q.status == 'rechazado')],
    ]
    for fila in filas_resumen:
        ws3.append(fila)
    ws3.column_dimensions['A'].width = 26
    ws3.column_dimensions['B'].width = 12

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"turnos_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return send_file(output, as_attachment=True, download_name=filename,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


# ─────────────────────────────────────────────
# PORTFOLIO
# ─────────────────────────────────────────────

@admin_bp.route('/portfolio')
def admin_portfolio():
    fotos = PortfolioImage.query.order_by(
        PortfolioImage.sort_order.asc(), PortfolioImage.id.asc()
    ).all()
    return render_template('admin_portfolio.html', fotos=fotos)


@admin_bp.route('/portfolio/subir', methods=['POST'])
def portfolio_upload():
    file_storage = request.files.get('foto')
    if not file_storage or not file_storage.filename:
        flash('Elegí una imagen para subir.', 'error')
        return redirect(url_for('admin.admin_portfolio'))
    if not images.allowed_file(file_storage.filename):
        flash('Formato no permitido (usá PNG, JPG o WebP).', 'error')
        return redirect(url_for('admin.admin_portfolio'))

    try:
        filename, thumb = images.save_image(
            file_storage, current_app.config['UPLOAD_FOLDER'], 'portfolio')
    except images.ImageError as e:
        flash(str(e), 'error')
        return redirect(url_for('admin.admin_portfolio'))

    ultimo = PortfolioImage.query.order_by(PortfolioImage.sort_order.desc()).first()
    foto = PortfolioImage(
        filename=filename,
        thumb_filename=thumb,
        title=request.form.get('title', '').strip() or None,
        category=request.form.get('category', '').strip() or None,
        alt_text=request.form.get('alt_text', '').strip() or None,
        sort_order=(ultimo.sort_order + 1) if ultimo else 0,
    )
    db.session.add(foto)
    db.session.commit()
    flash('Foto agregada al portfolio.', 'success')
    return redirect(url_for('admin.admin_portfolio'))


@admin_bp.route('/portfolio/<int:id>/actualizar', methods=['POST'])
def portfolio_update(id):
    foto = PortfolioImage.query.get_or_404(id)
    foto.title = request.form.get('title', '').strip() or None
    foto.category = request.form.get('category', '').strip() or None
    foto.alt_text = request.form.get('alt_text', '').strip() or None
    db.session.commit()
    flash('Foto actualizada.', 'success')
    return redirect(url_for('admin.admin_portfolio'))


@admin_bp.route('/portfolio/<int:id>/mover', methods=['POST'])
def portfolio_move(id):
    foto = PortfolioImage.query.get_or_404(id)
    direction = request.form.get('direction')
    vecina = None
    if direction == 'up':
        vecina = PortfolioImage.query.filter(
            PortfolioImage.sort_order < foto.sort_order
        ).order_by(PortfolioImage.sort_order.desc()).first()
    elif direction == 'down':
        vecina = PortfolioImage.query.filter(
            PortfolioImage.sort_order > foto.sort_order
        ).order_by(PortfolioImage.sort_order.asc()).first()

    if vecina is not None:
        foto.sort_order, vecina.sort_order = vecina.sort_order, foto.sort_order
        db.session.commit()
    return redirect(url_for('admin.admin_portfolio'))


@admin_bp.route('/portfolio/<int:id>/eliminar', methods=['POST'])
def portfolio_delete(id):
    foto = PortfolioImage.query.get_or_404(id)
    upload_root = current_app.config['UPLOAD_FOLDER']
    images.remove_image(upload_root, foto.filename)
    images.remove_image(upload_root, foto.thumb_filename)
    db.session.delete(foto)
    db.session.commit()
    flash('Foto eliminada.', 'info')
    return redirect(url_for('admin.admin_portfolio'))


# ─────────────────────────────────────────────
# PRESUPUESTOS
# ─────────────────────────────────────────────

@admin_bp.route('/quotes')
def admin_quotes():
    estado = request.args.get('estado', 'todos')
    query = QuoteRequest.query.order_by(QuoteRequest.created_at.desc())
    if estado != 'todos':
        query = query.filter_by(status=estado)
    quotes = query.all()
    stats = {
        'total': QuoteRequest.query.count(),
        'pendientes': QuoteRequest.query.filter_by(status='pendiente').count(),
        'aprobadas': QuoteRequest.query.filter_by(status='aprobado').count(),
        'agendadas': QuoteRequest.query.filter_by(status='agendado').count(),
        'rechazadas': QuoteRequest.query.filter_by(status='rechazado').count(),
    }
    return render_template('admin_quotes.html', quotes=quotes, stats=stats,
                           estado_filtro=estado)


@admin_bp.route('/quotes/<int:id>')
def admin_quote_detail(id):
    quote = QuoteRequest.query.get_or_404(id)
    client_wa = whatsapp.wa_me_link(
        phone=quote.telefono,
        text=f'Hola {quote.nombre}! Te escribo de Ink.Studio por tu presupuesto #{quote.id}.',
    )
    return render_template('admin_quote_detail.html', quote=quote, client_wa=client_wa)


@admin_bp.route('/quotes/<int:id>/estado', methods=['POST'])
def quote_cambiar_estado(id):
    quote = QuoteRequest.query.get_or_404(id)
    nuevo = request.form['estado']
    if nuevo in ['pendiente', 'aprobado', 'rechazado']:
        quote.status = nuevo
        db.session.commit()
        flash(f'Presupuesto marcado como "{nuevo}".', 'success')
    return redirect(url_for('admin.admin_quote_detail', id=quote.id))


@admin_bp.route('/quotes/<int:id>/eliminar', methods=['POST'])
def quote_delete(id):
    quote = QuoteRequest.query.get_or_404(id)
    upload_root = current_app.config['UPLOAD_FOLDER']
    for img in quote.images:
        images.remove_image(upload_root, img.filename)
    db.session.delete(quote)
    db.session.commit()
    flash('Presupuesto eliminado.', 'info')
    return redirect(url_for('admin.admin_quotes'))


# ─────────────────────────────────────────────
# CONFIGURACIÓN DEL ESTIMADOR
# ─────────────────────────────────────────────

@admin_bp.route('/estimator', methods=['GET', 'POST'])
def admin_estimator():
    if request.method == 'POST':
        action = request.form.get('action', 'save')
        if action == 'restore':
            seed_defaults(force=True)
            flash('Configuración restaurada a los valores por defecto.', 'info')
            return redirect(url_for('admin.admin_estimator'))

        updated = 0
        for group_key, fields in CONFIG_GROUPS:
            for key, _label in fields:
                raw = request.form.get(key, '').strip()
                if not raw:
                    continue
                try:
                    value = float(raw.replace(',', '.'))
                    if value < 0:
                        raise ValueError
                except ValueError:
                    flash(f'Valor inválido para "{key}": debe ser un número no negativo.', 'error')
                    return redirect(url_for('admin.admin_estimator'))
                set_config(key, value)
                updated += 1
        flash(f'Configuración guardada ({updated} valores).', 'success')
        return redirect(url_for('admin.admin_estimator'))

    cfg = get_config()
    return render_template('admin_estimator.html', config=cfg, groups=CONFIG_GROUPS)
